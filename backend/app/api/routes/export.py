"""Excel export endpoints — admin/director only."""

from __future__ import annotations

import io
from datetime import date, datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import AsyncSessionDep, CurrentUser, get_current_active_superuser
from app.models.org import Department, Role, UserCompanyRole
from app.models.project import Project
from app.models.task import Task
from app.models.user import User

router = APIRouter(prefix="/export", tags=["export"])

# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------

NAVY = "1F3864"
NAVY_LIGHT = "2E4F8A"
WHITE = "FFFFFF"
ROW_EVEN = "EBF0FA"
ROW_ODD = "FFFFFF"
ACCENT = "C6EFCE"
RED_LIGHT = "FFC7CE"
YELLOW_LIGHT = "FFEB9C"

_thin = Side(style="thin", color="CCCCCC")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _header_font() -> Font:
    return Font(name="Calibri", bold=True, color=WHITE, size=11)


def _header_fill() -> PatternFill:
    return PatternFill("solid", fgColor=NAVY)


def _sub_header_fill() -> PatternFill:
    return PatternFill("solid", fgColor=NAVY_LIGHT)


def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def _apply_header_row(ws: Any, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = _header_font()
        cell.fill = _header_fill()
        cell.alignment = _center()
        cell.border = _border


def _apply_data_row(ws: Any, row: int, values: list[Any], even: bool = False) -> None:
    fill = PatternFill("solid", fgColor=ROW_EVEN if even else ROW_ODD)
    for col, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.fill = fill
        cell.alignment = _left()
        cell.border = _border


def _auto_width(ws: Any, min_width: int = 12, max_width: int = 45) -> None:
    for col in ws.columns:
        length = max(
            len(str(cell.value or "")) for cell in col if cell.value is not None
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max_width, max(min_width, length + 4)
        )


def _fmt_dt(v: datetime | None) -> str:
    if v is None:
        return ""
    return v.strftime("%d/%m/%Y %H:%M")


def _fmt_d(v: date | None) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    return v.strftime("%d/%m/%Y")


# ---------------------------------------------------------------------------
# Auth helper — superuser OR director role
# ---------------------------------------------------------------------------

async def _require_admin_or_director(current_user: CurrentUser, session: AsyncSessionDep) -> User:
    if current_user.is_superuser:
        return current_user
    if current_user.company_id is None:
        raise HTTPException(status_code=403, detail="No company assigned")
    stmt = (
        select(UserCompanyRole)
        .join(Role, Role.id == UserCompanyRole.role_id)
        .where(
            UserCompanyRole.user_id == current_user.id,
            UserCompanyRole.company_id == current_user.company_id,
            Role.name == "director",
        )
    )
    result = await session.execute(stmt)
    if result.first() is None:
        raise HTTPException(status_code=403, detail="Director or superuser required")
    return current_user


AdminOrDirector = Depends(_require_admin_or_director)


# ---------------------------------------------------------------------------
# Data loaders
# ---------------------------------------------------------------------------

async def _load_users(session: AsyncSession, company_id: Any) -> list[Any]:
    stmt = (
        select(
            User.id,
            User.full_name,
            User.email,
            User.job_title,
            User.is_active,
            User.availability_status,
            User.created_at,
            Department.name.label("department_name"),
            Role.name.label("role_name"),
        )
        .outerjoin(Department, Department.id == User.department_id)
        .outerjoin(UserCompanyRole, UserCompanyRole.user_id == User.id)
        .outerjoin(Role, Role.id == UserCompanyRole.role_id)
        .where(User.company_id == company_id)
        .order_by(User.full_name)
    )
    result = await session.execute(stmt)
    return result.all()


async def _load_projects(session: AsyncSession, company_id: Any) -> list[Any]:
    stmt = (
        select(
            Project.id,
            Project.name,
            Project.code,
            Project.status,
            Project.start_date,
            Project.end_date,
            Project.description,
            User.full_name.label("pm_name"),
            Department.name.label("department_name"),
            func.count(Task.id).label("task_count"),
        )
        .outerjoin(User, User.id == Project.pm_id)
        .outerjoin(Department, Department.id == Project.department_id)
        .outerjoin(Task, (Task.project_id == Project.id) & (Task.is_deleted == False))  # noqa: E712
        .where(Project.company_id == company_id, Project.is_deleted == False)  # noqa: E712
        .group_by(Project.id, User.full_name, Department.name)
        .order_by(Project.code)
    )
    result = await session.execute(stmt)
    return result.all()


async def _load_tasks(session: AsyncSession, company_id: Any) -> list[Any]:
    assignee = User.__table__.alias("assignee")
    assignor = User.__table__.alias("assignor")
    stmt = (
        select(
            Task.id,
            Task.name,
            Task.priority,
            Task.status,
            Task.level,
            Task.start_time,
            Task.end_time,
            Task.actual_end_time,
            Task.module_tag,
            Task.created_at,
            Project.name.label("project_name"),
            Project.code.label("project_code"),
            assignee.c.full_name.label("assignee_name"),
            assignor.c.full_name.label("assignor_name"),
        )
        .join(Project, Project.id == Task.project_id)
        .outerjoin(assignee, assignee.c.id == Task.assignee_id)
        .outerjoin(assignor, assignor.c.id == Task.assignor_id)
        .where(
            Project.company_id == company_id,
            Task.is_deleted == False,  # noqa: E712
            Task.level <= 1,
        )
        .order_by(Project.code, Task.start_time)
    )
    result = await session.execute(stmt)
    return result.all()


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _build_analysis(wb: Workbook, users: list[Any], projects: list[Any], tasks: list[Any]) -> None:
    ws = wb.create_sheet("📊 Tổng Quan")

    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "BÁO CÁO PHÂN TÍCH HỆ THỐNG ERP"
    title_cell.font = Font(name="Calibri", bold=True, size=16, color=WHITE)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = _center()
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:H2")
    sub = ws["A2"]
    sub.value = f"Xuất ngày: {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')} UTC"
    sub.font = Font(name="Calibri", italic=True, size=10, color="555555")
    sub.alignment = _center()
    ws.row_dimensions[2].height = 18

    # --- KPI cards row 4-7 ---
    kpis = [
        ("Tổng nhân sự", len(users)),
        ("Nhân sự hoạt động", sum(1 for u in users if u.is_active)),
        ("Tổng dự án", len(projects)),
        ("Dự án đang chạy", sum(1 for p in projects if p.status == "in_progress")),
        ("Tổng task", len(tasks)),
        ("Task quá hạn", sum(
            1 for t in tasks
            if t.end_time and t.end_time.date() < date.today() and t.status not in ("done", "review")
        )),
    ]

    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 14

    for i, (label, value) in enumerate(kpis):
        col = i + 1
        label_cell = ws.cell(row=4, column=col, value=label)
        label_cell.font = Font(name="Calibri", bold=True, size=9, color="444444")
        label_cell.alignment = _center()
        label_cell.fill = PatternFill("solid", fgColor="F2F2F2")
        label_cell.border = _border

        value_cell = ws.cell(row=5, column=col, value=value)
        value_cell.font = Font(name="Calibri", bold=True, size=18, color=NAVY)
        value_cell.alignment = _center()
        value_cell.fill = PatternFill("solid", fgColor="DDEEFF")
        value_cell.border = _border

    # --- Task by status breakdown (row 8+) ---
    ws.cell(row=8, column=1, value="Phân tích Task theo Trạng thái").font = Font(bold=True, size=12, color=NAVY)
    ws.merge_cells("A8:D8")

    status_map: dict[str, int] = {}
    for t in tasks:
        status_map[t.status] = status_map.get(t.status, 0) + 1

    STATUS_LABELS = {
        "todo": "Chưa bắt đầu",
        "in_progress": "Đang làm",
        "review": "Đang review",
        "done": "Hoàn thành",
        "blocked": "Bị chặn",
    }

    _apply_header_row(ws, 9, ["Trạng thái", "Mã", "Số task", "Tỷ lệ %"])
    total_tasks = len(tasks) or 1
    for ri, (st, cnt) in enumerate(sorted(status_map.items())):
        _apply_data_row(ws, 10 + ri, [
            STATUS_LABELS.get(st, st),
            st,
            cnt,
            f"{cnt / total_tasks * 100:.1f}%",
        ], even=(ri % 2 == 0))

    # Pie chart for task status
    chart_start_row = 10
    chart_end_row = 10 + len(status_map) - 1
    pie = PieChart()
    pie.title = "Task theo trạng thái"
    pie.style = 10
    labels = Reference(ws, min_col=1, min_row=chart_start_row, max_row=chart_end_row)
    data = Reference(ws, min_col=3, min_row=chart_start_row - 1, max_row=chart_end_row)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.width = 14
    pie.height = 10
    ws.add_chart(pie, "F8")

    # --- Project by status (row after status table + 2) ---
    proj_row = 10 + len(status_map) + 2
    ws.cell(row=proj_row, column=1, value="Phân tích Dự án theo Trạng thái").font = Font(bold=True, size=12, color=NAVY)
    ws.merge_cells(f"A{proj_row}:D{proj_row}")

    proj_status_map: dict[str, int] = {}
    for p in projects:
        proj_status_map[p.status] = proj_status_map.get(p.status, 0) + 1

    _apply_header_row(ws, proj_row + 1, ["Trạng thái", "Số dự án", "Task / dự án TB"])
    for ri, (st, cnt) in enumerate(sorted(proj_status_map.items())):
        proj_tasks = sum(p.task_count for p in projects if p.status == st)
        avg = round(proj_tasks / cnt, 1) if cnt else 0
        _apply_data_row(ws, proj_row + 2 + ri, [st, cnt, avg], even=(ri % 2 == 0))

    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.freeze_panes = "A3"


def _build_accounts(wb: Workbook, users: list[Any]) -> None:
    ws = wb.create_sheet("👥 Nhân Sự & Vai Trò")

    headers = [
        "Họ tên", "Email", "Chức vụ", "Phòng ban",
        "Vai trò hệ thống", "Trạng thái", "Hoạt động", "Ngày tạo",
    ]
    _apply_header_row(ws, 1, headers)

    # Deduplicate users (multiple rows due to role join)
    seen: set[Any] = set()
    data_rows: list[list[Any]] = []
    for u in users:
        if u.id in seen:
            continue
        seen.add(u.id)
        data_rows.append([
            u.full_name or "",
            u.email,
            u.job_title or "",
            u.department_name or "",
            u.role_name or "",
            "Hoạt động" if u.is_active else "Vô hiệu",
            u.availability_status or "",
            _fmt_dt(u.created_at),
        ])

    for ri, row in enumerate(data_rows):
        _apply_data_row(ws, ri + 2, row, even=(ri % 2 == 0))
        # Color inactive users
        if row[5] == "Vô hiệu":
            for col in range(1, 9):
                ws.cell(row=ri + 2, column=col).fill = PatternFill("solid", fgColor=RED_LIGHT)

    _auto_width(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(data_rows) + 1}"


def _build_tasks(wb: Workbook, tasks: list[Any]) -> None:
    ws = wb.create_sheet("✅ Task")

    headers = [
        "Dự án", "Mã DA", "Tên task", "Cấp",
        "Ưu tiên", "Trạng thái", "Người giao", "Người nhận",
        "Bắt đầu", "Kết thúc", "Kết thúc thực tế", "Module",
    ]
    _apply_header_row(ws, 1, headers)

    today = date.today()
    for ri, t in enumerate(tasks):
        overdue = (
            t.end_time and t.end_time.date() < today
            and t.status not in ("done", "review")
        )
        priority_flag = {"high": "🔴", "medium": "🟡", "low": "🟢"}.get(t.priority, "")
        row = [
            t.project_name,
            t.project_code,
            t.name,
            {0: "Đầu việc", 1: "Task", 2: "Sub-task"}.get(t.level, f"L{t.level}"),
            f"{priority_flag} {t.priority}",
            t.status,
            t.assignor_name or "",
            t.assignee_name or "",
            _fmt_dt(t.start_time),
            _fmt_dt(t.end_time),
            _fmt_dt(t.actual_end_time),
            t.module_tag or "",
        ]
        _apply_data_row(ws, ri + 2, row, even=(ri % 2 == 0))
        if overdue:
            for col in range(1, 13):
                ws.cell(row=ri + 2, column=col).fill = PatternFill("solid", fgColor=RED_LIGHT)
        elif t.status == "done":
            for col in range(1, 13):
                ws.cell(row=ri + 2, column=col).fill = PatternFill("solid", fgColor=ACCENT)

    _auto_width(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{len(tasks) + 1}"


def _build_projects(wb: Workbook, projects: list[Any]) -> None:
    ws = wb.create_sheet("📁 Dự Án")

    headers = [
        "Mã", "Tên dự án", "Trạng thái", "Quản lý dự án",
        "Phòng ban", "Bắt đầu", "Kết thúc", "Số task", "Mô tả",
    ]
    _apply_header_row(ws, 1, headers)

    today = date.today()
    for ri, p in enumerate(projects):
        overdue = p.end_date and p.end_date < today and p.status not in ("done", "completed")
        row = [
            p.code,
            p.name,
            p.status,
            p.pm_name or "",
            p.department_name or "",
            _fmt_d(p.start_date),
            _fmt_d(p.end_date),
            p.task_count,
            (p.description or "")[:200],
        ]
        _apply_data_row(ws, ri + 2, row, even=(ri % 2 == 0))
        if overdue:
            for col in range(1, 10):
                ws.cell(row=ri + 2, column=col).fill = PatternFill("solid", fgColor=RED_LIGHT)

    _auto_width(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(projects) + 1}"


def _build_monthly(wb: Workbook, tasks: list[Any]) -> None:
    ws = wb.create_sheet("📈 Báo Cáo Tháng")

    # Monthly task completion count
    monthly: dict[str, dict[str, int]] = {}
    for t in tasks:
        if t.created_at:
            key = t.created_at.strftime("%Y-%m")
            if key not in monthly:
                monthly[key] = {"created": 0, "done": 0}
            monthly[key]["created"] += 1
            if t.status in ("done", "review"):
                monthly[key]["done"] += 1

    headers = ["Tháng", "Task tạo mới", "Task hoàn thành", "Tỷ lệ hoàn thành %"]
    _apply_header_row(ws, 1, headers)

    sorted_months = sorted(monthly.keys())
    for ri, month in enumerate(sorted_months):
        created = monthly[month]["created"]
        done = monthly[month]["done"]
        rate = f"{done / created * 100:.1f}%" if created else "0%"
        _apply_data_row(ws, ri + 2, [month, created, done, rate], even=(ri % 2 == 0))

    # Bar chart
    if sorted_months:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Task tạo mới vs Hoàn thành theo tháng"
        chart.y_axis.title = "Số task"
        chart.x_axis.title = "Tháng"

        data_ref = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=len(sorted_months) + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=len(sorted_months) + 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 22
        chart.height = 14
        ws.add_chart(chart, "F2")

    _auto_width(ws)
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Endpoint
# ---------------------------------------------------------------------------

def _make_filename(prefix: str) -> str:
    return f"{prefix}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.xlsx"


async def _generate_workbook(session: AsyncSession, company_id: Any) -> Workbook:
    users, projects, tasks = (
        await _load_users(session, company_id),
        await _load_projects(session, company_id),
        await _load_tasks(session, company_id),
    )

    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    _build_analysis(wb, users, projects, tasks)
    _build_accounts(wb, users)
    _build_tasks(wb, tasks)
    _build_projects(wb, projects)
    _build_monthly(wb, tasks)

    return wb


def _stream_wb(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/full")
async def export_full(
    session: AsyncSessionDep,
    current_user: CurrentUser,
    _: User = AdminOrDirector,
) -> StreamingResponse:
    """Export toàn bộ dữ liệu: nhân sự, dự án, task, phân tích."""
    if current_user.company_id is None and not current_user.is_superuser:
        raise HTTPException(status_code=400, detail="No company")
    company_id = current_user.company_id
    wb = await _generate_workbook(session, company_id)
    return _stream_wb(wb, _make_filename("saree_export_full"))


@router.get("/accounts")
async def export_accounts(
    session: AsyncSessionDep,
    current_user: CurrentUser,
    _: User = AdminOrDirector,
) -> StreamingResponse:
    """Export danh sách nhân sự và vai trò."""
    company_id = current_user.company_id
    users = await _load_users(session, company_id)

    wb = Workbook()
    wb.remove(wb.active)
    _build_accounts(wb, users)
    _build_analysis(wb, users, [], [])

    return _stream_wb(wb, _make_filename("saree_export_accounts"))


@router.get("/tasks")
async def export_tasks(
    session: AsyncSessionDep,
    current_user: CurrentUser,
    _: User = AdminOrDirector,
) -> StreamingResponse:
    """Export danh sách task và dự án."""
    company_id = current_user.company_id
    projects = await _load_projects(session, company_id)
    tasks = await _load_tasks(session, company_id)

    wb = Workbook()
    wb.remove(wb.active)
    _build_tasks(wb, tasks)
    _build_projects(wb, projects)
    _build_monthly(wb, tasks)

    return _stream_wb(wb, _make_filename("saree_export_tasks"))
